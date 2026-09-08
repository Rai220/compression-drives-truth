# -*- coding: utf-8 -*-
"""GigaRnD deck: 'Compression Favors Consistency, Not Truth'. Russian, 16:9.
Opens with a live empirical probe of public models (worldview_results.json),
then the controlled experiments, then back to the hook. python-pptx."""
import json, os, collections
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
WV = json.load(open(os.path.join(HERE, "worldview_results.json"), encoding="utf-8"))
def _try(name):
    p=os.path.join(HERE,name)
    return json.load(open(p,encoding="utf-8")) if os.path.exists(p) else None
TCM = _try("tcm_results.json")
BASE = _try("base_probe2_results.json")      # length-matched (headline)
BASE_LEN = _try("base_probe_results.json")   # length-confounded (caveat illustration)

# ---- palette ----
INK=RGBColor(0x1A,0x1A,0x1A); GRAPHITE=RGBColor(0x33,0x3A,0x40)
GREEN=RGBColor(0x18,0x9E,0x3B); GREEN_D=RGBColor(0x0F,0x6E,0x2A)
MUTED=RGBColor(0x6B,0x72,0x78); LIGHT=RGBColor(0xF2,0xF4,0xF3)
BAND=RGBColor(0xE8,0xF3,0xEC); WHITE=RGBColor(0xFF,0xFF,0xFF)
RED=RGBColor(0xC0,0x39,0x2B); AMBER=RGBColor(0xB9,0x7A,0x0E); BLUE=RGBColor(0x2C,0x5A,0x8C)
REDBG=RGBColor(0xFB,0xEC,0xEA); AMBBG=RGBColor(0xFB,0xF3,0xDE)
FONT="Calibri"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def box(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    return tb,tf
def para(tf,t,sz,color=INK,bold=False,first=False,align=PP_ALIGN.LEFT,sa=6,it=False,font=FONT):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(sa); p.space_before=Pt(0)
    r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold; r.font.italic=it
    r.font.name=font; r.font.color.rgb=color; return p
def rect(s,x,y,w,h,fill,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h); sh.fill.solid()
    sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(0.75)
    sh.shadow.inherit=False; return sh
def bullet(tf,t,sz=15,color=INK,bold=False,first=False,sa=9,mark="—",it=False):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after=Pt(sa); p.space_before=Pt(0)
    r1=p.add_run(); r1.text=mark+"  "; r1.font.size=Pt(sz); r1.font.name=FONT
    r1.font.bold=True; r1.font.color.rgb=GREEN
    r2=p.add_run(); r2.text=t; r2.font.size=Pt(sz); r2.font.name=FONT
    r2.font.bold=bold; r2.font.italic=it; r2.font.color.rgb=color; return p
def header(s,kicker,title):
    rect(s,0,0,Inches(0.22),Inches(7.5),GREEN)
    _,tf=box(s,Inches(0.6),Inches(0.4),Inches(12.2),Inches(0.4))
    para(tf,kicker.upper(),12.5,GREEN_D,bold=True,first=True,sa=2)
    _,tf2=box(s,Inches(0.6),Inches(0.76),Inches(12.2),Inches(1.0))
    para(tf2,title,26,INK,bold=True,first=True,sa=0)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.62),Inches(1.58),Inches(1.1),Pt(3))
    r.fill.solid(); r.fill.fore_color.rgb=GREEN; r.line.fill.background(); r.shadow.inherit=False
def footer(s,n=None):
    n=len(prs.slides._sldIdLst)  # auto-number by real slide position
    _,tf=box(s,Inches(0.6),Inches(7.04),Inches(9),Inches(0.34))
    para(tf,"Сжатие предпочитает согласованность, а не истину  ·  GigaRnD",9,MUTED,first=True)
    _,tf2=box(s,Inches(11.9),Inches(7.04),Inches(0.9),Inches(0.34))
    para(tf2,str(n),9,MUTED,first=True,align=PP_ALIGN.RIGHT)
def set_cell(cell,t,sz=13,color=INK,bold=False,fill=WHITE,align=PP_ALIGN.CENTER):
    cell.margin_left=Inches(0.06); cell.margin_right=Inches(0.06)
    cell.margin_top=Inches(0.03); cell.margin_bottom=Inches(0.03)
    cell.vertical_anchor=MSO_ANCHOR.MIDDLE; cell.fill.solid(); cell.fill.fore_color.rgb=fill
    tf=cell.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold; r.font.name=FONT
    r.font.color.rgb=color
def style_table(tbl):
    tbl._tbl.tblPr.set('firstRow','0'); tbl._tbl.tblPr.set('bandRow','0')

VCOL={"Да":BLUE,"Нет":GREEN_D,"Спорно":AMBER,"—":MUTED,"?":MUTED}
VBG ={"Да":RGBColor(0xEA,0xF0,0xF7),"Нет":BAND,"Спорно":AMBBG,"—":LIGHT,"?":LIGHT}
MODEL_ORDER=["GPT-4o-mini","DeepSeek-V3","Qwen2.5-72B","Llama-3.3-70B","Mistral-Large","GLM-5.2"]
QLABEL={"astro":"Астрология","homeo":"Гомеопатия","gmo":"ГМО безопасны?",
        "freewill":"Свобода воли","morality":"Объектив. мораль","aiconsc":"Сознание ИИ"}
QORDER=["astro","homeo","gmo","freewill","morality","aiconsc"]

def verdicts(qid):
    a=WV[qid]["answers"]; return {m:a.get(m,{}).get("verdict","—") for m in MODEL_ORDER}
def tally(qid):
    c=collections.Counter(verdicts(qid).values());
    top=c.most_common(1)[0]; return top[0],top[1],len(MODEL_ORDER)
def reason(qid,model):
    return WV[qid]["answers"].get(model,{}).get("reason","")

# =====================================================================
# 1 — TITLE
# =====================================================================
s=slide(); rect(s,0,0,Inches(13.333),Inches(7.5),WHITE)
rect(s,0,0,Inches(0.32),Inches(7.5),GREEN); rect(s,0,Inches(5.5),Inches(13.333),Inches(0.04),BAND)
_,tf=box(s,Inches(0.9),Inches(1.55),Inches(11.7),Inches(2.4))
para(tf,"Сжатие предпочитает согласованность,\nа не истину",39,INK,bold=True,first=True,sa=10)
_,tf2=box(s,Inches(0.92),Inches(3.6),Inches(11.4),Inches(1.5))
para(tf2,"Как языковые модели выбирают «во что верить» — и почему",19,GRAPHITE,first=True,sa=4)
para(tf2,"согласованная ложь может обойти фильтр правды",19,GRAPHITE)
_,tf3=box(s,Inches(0.92),Inches(5.72),Inches(11.4),Inches(1.3))
para(tf3,"Внутренняя конференция GigaRnD",15,GREEN_D,bold=True,first=True,sa=3)
para(tf3,"От эмпирики публичных моделей — к контролируемому эксперименту 3.5M–3.08B  ·  NeurIPS 2026",13,MUTED)

# =====================================================================
# 2 — HOOK: we asked 6 models
# =====================================================================
s=slide(); header(s,"Эмпирический хук","6 моделей от разных команд — почти одна картина мира")
# grid table: rows = questions, cols = models
nrow=len(QORDER)+1; ncol=len(MODEL_ORDER)+1
t=s.shapes.add_table(nrow,ncol,Inches(0.6),Inches(1.85),Inches(12.15),Inches(4.35)).table
style_table(t)
t.columns[0].width=Inches(2.55)
for c in range(1,ncol): t.columns[c].width=Inches(1.6)
set_cell(t.cell(0,0),"вопрос",12,WHITE,bold=True,fill=GRAPHITE,align=PP_ALIGN.LEFT)
for j,m in enumerate(MODEL_ORDER): set_cell(t.cell(0,j+1),m,11,WHITE,bold=True,fill=GRAPHITE)
for i,qid in enumerate(QORDER):
    set_cell(t.cell(i+1,0),QLABEL[qid],12,INK,bold=True,fill=LIGHT,align=PP_ALIGN.LEFT)
    vd=verdicts(qid)
    for j,m in enumerate(MODEL_ORDER):
        v=vd[m]; set_cell(t.cell(i+1,j+1),v,12,VCOL.get(v,INK),bold=True,fill=VBG.get(v,WHITE))
_,tc=box(s,Inches(0.6),Inches(6.4),Inches(12.1),Inches(0.55))
para(tc,"Разные команды, разные данные — а карта одна: твёрдый скепсис к псевдонауке, "
        "единодушное «Спорно» на философских вопросах. Расхождение лишь на краях (Qwen · астрология). "
        "Ответы вживую через OpenRouter, temp=0.",12,MUTED,first=True)
footer(s,2)

# =====================================================================
# 3 — THE PARADOX (astrology)
# =====================================================================
s=slide(); header(s,"Парадокс","Астрологических текстов — море. Модели говорят «Нет».")
av,ac,an=tally("astro")
rect(s,Inches(0.6),Inches(1.9),Inches(5.7),Inches(2.0),BAND)
_,tf=box(s,Inches(0.85),Inches(2.05),Inches(5.2),Inches(1.7),anchor=MSO_ANCHOR.MIDDLE)
para(tf,f"{ac} из {an} моделей: «{av}»",30,GREEN_D,bold=True,first=True,sa=6)
para(tf,"на вопрос, влияет ли положение звёзд при рождении на судьбу.",14,INK)
_,tf2=box(s,Inches(0.6),Inches(4.15),Inches(5.9),Inches(2.6))
bullet(tf2,"«Больше текста» ≠ «больше веры»: астрологии в данных огромно, "
           "но модели её отвергают.",15,first=True,bold=True)
bullet(tf2,"Значит, дело не в частоте. И не в явной «правде» — базовый объектив "
           "её не знает.",15)
bullet(tf2,"Так КАК модель выбирает свою точку зрения?",16,color=RED,bold=True)
# actual model quotes
rect(s,Inches(6.75),Inches(1.9),Inches(5.95),Inches(4.85),LIGHT)
_,tq=box(s,Inches(7.0),Inches(2.05),Inches(5.5),Inches(4.6))
para(tq,"Что реально ответили модели:",13.5,GREEN_D,bold=True,first=True,sa=9)
shown=0
for m in MODEL_ORDER:
    r=reason("astro",m)
    if r and shown<4:
        para(tq,f"{m}",12,GRAPHITE,bold=True,sa=1)
        para(tq,f"«{r}»",12.5,INK,it=True,sa=9)
        shown+=1
footer(s,3)

# =====================================================================
# 4 — THREE NAIVE ANSWERS
# =====================================================================
s=slide(); header(s,"Гипотезы","Три «очевидных» объяснения — и почему их мало")
cards=[("Частота","Чаще встречается → в это и верит.",
        "Не работает: астрологии много, но её отвергают. И наш эксперимент держит частоту 50/50.",RED),
       ("Истина","Объектив «знает» правду.",
        "Не работает: при равной частоте объектив НЕ отличает согласованную ложь от правды (покажем).",RED),
       ("Алайнмент","RLHF/политика прививают скепсис.",
        "Часть ответа — но расхождение видно и в БАЗЕ, до алайнмента. Значит есть более глубокий слой.",AMBER)]
x=Inches(0.6)
for title,claim,verdict,col in cards:
    rect(s,x,Inches(1.95),Inches(3.95),Inches(4.4),LIGHT)
    rect(s,x,Inches(1.95),Inches(3.95),Inches(0.6),GRAPHITE)
    _,th=box(s,x+Inches(0.22),Inches(2.0),Inches(3.5),Inches(0.5),anchor=MSO_ANCHOR.MIDDLE)
    para(th,title,16,WHITE,bold=True,first=True)
    _,tb=box(s,x+Inches(0.22),Inches(2.75),Inches(3.55),Inches(3.5))
    para(tb,claim,14,INK,bold=True,first=True,sa=10)
    para(tb,"✕  "+verdict if col==RED else "≈  "+verdict,13,col,sa=0)
    x=Inches(x.inches+4.12)
_,tf=box(s,Inches(0.6),Inches(6.5),Inches(12),Inches(0.5))
para(tf,"Разберём глубокий слой на чистом эксперименте, где частота, формат и домен зафиксированы.",
     14.5,GREEN_D,bold=True,first=True)
footer(s,4)

# =====================================================================
# 4b — ALIGNMENT: TCM East/West split
# =====================================================================
if TCM:
    s=slide(); header(s,"Слой 1 — алайнмент","Китайская медицина: Восток и Запад расходятся")
    ta=TCM["tcm"]["answers"]
    west=[(m,d) for m,d in ta.items() if d.get("origin","").startswith("West")]
    china=[(m,d) for m,d in ta.items() if d.get("origin")=="China"]
    def panel(x,title,items,col):
        rect(s,x,Inches(1.95),Inches(5.9),Inches(3.15),LIGHT)
        rect(s,x,Inches(1.95),Inches(5.9),Inches(0.55),col)
        _,t=box(s,x+Inches(0.25),Inches(2.0),Inches(5.4),Inches(0.45),anchor=MSO_ANCHOR.MIDDLE)
        para(t,title,15,WHITE,bold=True,first=True)
        _,b=box(s,x+Inches(0.25),Inches(2.68),Inches(5.4),Inches(2.4))
        for i,(m,d) in enumerate(items):
            v=d["verdict"]; p=b.paragraphs[0] if i==0 else b.add_paragraph(); p.space_after=Pt(5)
            r1=p.add_run(); r1.text=f"{m}: "; r1.font.size=Pt(13.5); r1.font.name=FONT; r1.font.color.rgb=INK
            r2=p.add_run(); r2.text=v; r2.font.size=Pt(13.5); r2.font.bold=True; r2.font.name=FONT
            r2.font.color.rgb=VCOL.get(v,INK)
    panel(Inches(0.6),"Запад (US / EU)",west,BLUE)
    panel(Inches(6.85),"Китай",china,RED)
    rect(s,Inches(0.6),Inches(5.35),Inches(12.15),Inches(1.35),AMBBG)
    _,tt=box(s,Inches(0.9),Inches(5.45),Inches(11.5),Inches(1.2),anchor=MSO_ANCHOR.MIDDLE)
    para(tt,"Западные модели готовы сказать «Нет». Китайские — только «Спорно» (или молчат).",
         18,INK,bold=True,first=True,sa=4)
    para(tt,"Одна и та же наука — разная политика согласования. Алайнмент сдвигает картину мира. "
            "Но расхождение видно и глубже — значит, есть слой под ним.",13.5,GRAPHITE)
    footer(s)

# =====================================================================
# 4c — BASE MODEL PROBE (rendered once data is available)
# =====================================================================
if BASE:
    s=slide(); header(s,"Слой 2 — сам претрейн","А «чистая» base-модель без RLHF? Нет компаса истины")
    _,tf=box(s,Inches(0.6),Inches(1.78),Inches(4.75),Inches(4.9))
    para(tf,"Base-модель не следует инструкциям — «спросить» нельзя. Меряем "
            "ПРАВДОПОДОБИЕМ (share), как в статье: P(научный вердикт) против "
            "P(псевдонаучного), длина контролируется.",13,INK,first=True,sa=9)
    bullet(tf,"На КАЖДОМ псевдонаучном утверждении share НИЖЕ 50 — база "
              "склоняется к credulous-версии. Претрейн НЕ отвергает псевдонауку.",
              13,color=RED,bold=True,mark="•")
    bullet(tf,"Контроль (эволюция) — 77–84: проба ловит реальный наклон, значит "
              "низкие клетки не шум.",13,mark="•")
    bullet(tf,"Скептицизм из слайда 2 — это АЛАЙНМЕНТ, а не претрейн. У объектива "
              "претрейна нет компаса истины — ровно тезис статьи.",13,color=GREEN_D,
              bold=True,mark="•")
    repos=list(BASE.keys()); topics=list(BASE[repos[0]].keys())
    def shortrepo(r): return r.split("/")[-1].replace("Qwen2.5-","Qwen ")
    t=s.shapes.add_table(len(topics)+1,len(repos)+1,Inches(5.55),Inches(1.8),
                         Inches(7.15),Inches(4.5)).table
    style_table(t); t.columns[0].width=Inches(3.35)
    for c in range(1,len(repos)+1): t.columns[c].width=Inches(1.27)
    set_cell(t.cell(0,0),"Тема — share правды, %",11.5,WHITE,bold=True,fill=GRAPHITE,align=PP_ALIGN.LEFT)
    for j,r in enumerate(repos): set_cell(t.cell(0,j+1),shortrepo(r),11.5,WHITE,bold=True,fill=GRAPHITE)
    for i,tp in enumerate(topics):
        ctrl="+" in tp
        set_cell(t.cell(i+1,0),tp,11.5,INK,bold=True,fill=BAND if ctrl else LIGHT,align=PP_ALIGN.LEFT)
        for j,r in enumerate(repos):
            sh=BASE[r][tp]["share"]*100
            col=GREEN_D if sh>=55 else (RED if sh<=45 else AMBER)
            bg=BAND if sh>=55 else (REDBG if sh<=45 else AMBBG)
            set_cell(t.cell(i+1,j+1),f"{sh:.0f}",12.5,col,bold=True,fill=bg)
    _,tc=box(s,Inches(5.55),Inches(6.42),Inches(7.15),Inches(0.6))
    para(tc,"Зелёный = склон к науке (>55), красный = к псевдонауке (<45). Base Qwen2.5, "
            "длина ответов уравнена. В наивной (неуравненной) пробе знак share двигает даже "
            "частица «не» — тот же length-confound, что вскрывает статья.",10.5,MUTED,first=True)
    footer(s)

# =====================================================================
# 5 — SETUP
# =====================================================================
s=slide(); header(s,"Постановка","Контролируемый эксперимент")
_,tf=box(s,Inches(0.6),Inches(1.95),Inches(6.0),Inches(3.0))
bullet(tf,"Обучаем трансформеры с нуля на мат-корпусах.",15.5,first=True)
bullet(tf,"Каждая задача встречается и с верным, и с неверным решением — "
          "частота, формат, домен фиксированы.",15.5)
bullet(tf,"Варьируем ТОЛЬКО структуру ошибки.",15.5,bold=True)
def cond_card(x,title,sub,col):
    rect(s,x,Inches(1.95),Inches(3.05),Inches(4.5),LIGHT)
    rect(s,x,Inches(1.95),Inches(3.05),Inches(0.6),col)
    _,t=box(s,x+Inches(0.2),Inches(2.0),Inches(2.7),Inches(0.5),anchor=MSO_ANCHOR.MIDDLE)
    para(t,title,15,WHITE,bold=True,first=True)
    _,b=box(s,x+Inches(0.2),Inches(2.75),Inches(2.7),Inches(3.5))
    for i,l in enumerate(sub): para(b,l,12.5,INK,first=(i==0),sa=7)
cond_card(Inches(6.75),"random",["Ошибка случайна.","Много ложных форм.",
    "Модель не знает, какая выпадет.","","→ непредсказуемая ложь"],GRAPHITE)
cond_card(Inches(9.95),"coherent",["Одно правило на все задачи.","Одна ложная строка.",
    "Предсказуемо, как истина.","","→ согласованная ложь"],GREEN_D)
_,tf2=box(s,Inches(0.6),Inches(5.15),Inches(6.0),Inches(1.2))
para(tf2,"Ключевой аналог астрологии — это coherent: внутренне "
         "согласованная, но ложная система.",14.5,RED,bold=True,first=True)
footer(s,5)

# =====================================================================
# 6 — MASS CONTROL
# =====================================================================
s=slide(); header(s,"Правка метрики","Pair accuracy обманывает — меряем массу")
_,tf=box(s,Inches(0.6),Inches(1.9),Inches(6.1),Inches(3.0))
bullet(tf,"Pair accuracy сравнивает верное с ОДНОЙ ложной строкой.",15,first=True)
bullet(tf,"При N ложных формах модель, просто повторяющая частоты, выигрывает "
          "попарно, ничего не предпочитая.",15)
bullet(tf,"Честная мера — aggregate share: масса верного против ВСЕГО ложного "
          "семейства. Null ровно 0.5.",15,bold=True)
_,tff=box(s,Inches(0.6),Inches(4.55),Inches(6.1),Inches(0.8))
para(tff,"share = P(верно) / [ P(верно) + Σ P(ложь_i) ]",17,INK,bold=True,first=True,font="Consolas")
rows=[("Мера (86M, random)","Значение","над 0.5"),("pair accuracy","87.4 %","+37 пп"),
      ("realized (1 строка)","77.6 %","+27 пп"),("aggregate share","59.0 %","+9 пп")]
gt=s.shapes.add_table(4,3,Inches(7.0),Inches(2.15),Inches(5.7),Inches(2.4)).table
style_table(gt); gt.columns[0].width=Inches(2.9); gt.columns[1].width=Inches(1.4); gt.columns[2].width=Inches(1.4)
for j,h in enumerate(rows[0]): set_cell(gt.cell(0,j),h,13,WHITE,bold=True,fill=GRAPHITE,align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
for i in range(1,4):
    hl=(i==3)
    for j,v in enumerate(rows[i]):
        set_cell(gt.cell(i,j),v,13.5,GREEN_D if hl else INK,bold=hl,fill=BAND if hl else WHITE,
                 align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
_,tc=box(s,Inches(7.0),Inches(4.9),Inches(5.7),Inches(1.4))
para(tc,"Попарные метрики переоценивают эффект ~в 4 раза — и меняют знак ниже 86M.",14.5,RED,bold=True,first=True)
footer(s,6)

# =====================================================================
# 7 — SURVIVED / WITHDRAWN
# =====================================================================
s=slide(); header(s,"Итог контролей (авг 2026)","Что выжило и что отозвано")
rect(s,Inches(0.6),Inches(1.95),Inches(5.9),Inches(4.5),BAND); rect(s,Inches(0.6),Inches(1.95),Inches(5.9),Inches(0.55),GREEN)
_,t=box(s,Inches(0.85),Inches(2.0),Inches(5.4),Inches(0.45),anchor=MSO_ANCHOR.MIDDLE)
para(t,"✓  ВЫЖИЛО",15,WHITE,bold=True,first=True)
_,b=box(s,Inches(0.85),Inches(2.68),Inches(5.4),Inches(3.6))
bullet(b,"86M: верное держит 59.0 %±0.8 массы против random (≈+9 пп).",14,first=True)
bullet(b,"Coherent НЕ фильтруется ни на каком размере — робастная половина.",14)
bullet(b,"Реплика на независимом корпусе: 59.4 %.",14)
rect(s,Inches(6.85),Inches(1.95),Inches(5.85),Inches(4.5),LIGHT); rect(s,Inches(6.85),Inches(1.95),Inches(5.85),Inches(0.55),RED)
_,t2=box(s,Inches(7.1),Inches(2.0),Inches(5.3),Inches(0.45),anchor=MSO_ANCHOR.MIDDLE)
para(t2,"✕  ОТОЗВАНО",15,WHITE,bold=True,first=True)
_,b2=box(s,Inches(7.1),Inches(2.68),Inches(5.3),Inches(3.6))
bullet(b2,"Multi-rule: pair acc 47→79, а масса 47.9→47.8. Это mass splitting.",14,first=True,mark="—")
bullet(b2,"Description-length ladder: масса плоская (55.6/51.9/52.6/51.3).",14,mark="—")
bullet(b2,"«Coherent ниже случайности» — это длина завершения.",14,mark="—")
footer(s,7)

# =====================================================================
# 8 — NEW: learnable selector
# =====================================================================
s=slide(); header(s,"Новый результат · главное","Learnable selector: знак предпочтения переключается")
_,tf=box(s,Inches(0.6),Inches(1.78),Inches(4.5),Inches(4.8))
para(tf,"Держим ложное семейство фиксированным (10 строк). Меняем ТОЛЬКО "
        "выучиваемость селектора.",14.5,INK,first=True,sa=10)
bullet(tf,"невыучиваем → предпочитается ВЕРНОЕ",13.5,mark="•")
bullet(tf,"выучиваем → предпочитается ЛОЖНОЕ",13.5,bold=True,mark="•",color=RED)
bullet(tf,"частично → ровно на нуле",13.5,mark="•")
para(tf,"Причина — предсказуемость, а не длина описания. Число мод постоянно.",
     13.5,GREEN_D,bold=True,sa=0)
data=[("Селектор (474M)","share","real.","знак"),("random (непредск.)","54.9","≈92","+5"),
      ("coherent (1 правило)","49.9","—","0"),("learnable (1 цифра)","43.4","66","−7"),
      ("k2 (сумма 2 цифр)","49.2","89","0"),("k3 (сумма 3 цифр)","48.8","92","0"),
      ("learnable_mod (все)","55.4","93","+5")]
t=s.shapes.add_table(7,4,Inches(5.35),Inches(1.72),Inches(7.35),Inches(4.6)).table
style_table(t)
for w,c in zip((3.05,1.4,1.35,1.2),range(4)): t.columns[c].width=Inches(w)
for j,h in enumerate(data[0]): set_cell(t.cell(0,j),h,12.5,WHITE,bold=True,fill=GRAPHITE,align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
sign={"+5":GREEN_D,"0":MUTED,"−7":RED}
for i in range(1,7):
    learn=data[i][0].startswith("learnable (")
    for j,v in enumerate(data[i]):
        col=INK;bold=False;fill=WHITE
        if learn: fill=REDBG; bold=True
        if j==3: col=sign.get(v,INK); bold=True
        set_cell(t.cell(i,j),v,12.5,col,bold=bold,fill=fill,align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
_,tc=box(s,Inches(5.35),Inches(6.45),Inches(7.35),Inches(0.5))
para(tc,"family share, length-matched, null 50.  418M: random +6.8 / learnable −5.3.",11.5,MUTED,first=True)
footer(s,8)

# =====================================================================
# 9 — MECHANISM
# =====================================================================
s=slide(); header(s,"Механизм","realized share — медиатор; знак не артефакт")
_,tf=box(s,Inches(0.6),Inches(2.0),Inches(6.0),Inches(4.4))
bullet(tf,"Выученный селектор концентрирует ложную массу на предсказуемой "
          "строке → модель «коммитится» в ложь → share ниже нуля.",15,first=True)
bullet(tf,"Градиент learnable→k2→k3: realized 66→89→92, share 43→49→49. Монотонно.",15)
bullet(tf,"learnable_mod детерминирован, но НЕ выучивается → читается как random. "
          "Важна реальная выучиваемость, не детерминизм.",15)
rect(s,Inches(6.9),Inches(2.0),Inches(5.8),Inches(4.2),BAND)
_,t=box(s,Inches(7.15),Inches(2.2),Inches(5.35),Inches(3.9))
para(t,"Почему below-null = предпочтение, а не некомпетентность?",15,GREEN_D,bold=True,first=True,sa=8)
para(t,"random и learnable обучены на ОДНИХ верных выводах — считают правильный "
        "ответ одинаково.",14,INK,sa=8)
para(t,"random выше нуля, learnable ниже — при равной компетентности. Причина — "
        "именно селектор.",14,INK,sa=8)
para(t,"Это causal identification, которую раньше «не знали, как построить».",14,INK,bold=True,it=True)
footer(s,9)

# =====================================================================
# 10 — SCALE
# =====================================================================
s=slide(); header(s,"Масштаб","Контраст держится до настоящих 3.08B")
_,tf=box(s,Inches(0.6),Inches(1.95),Inches(4.4),Inches(4.4))
bullet(tf,"Раньше потолок 490M. Теперь — Qwen3 на 1.12B / 2.16B / 3.08B (2 сида).",15,first=True)
bullet(tf,"random выше нуля, coherent на нуле, learnable ниже нуля — на каждом размере.",15)
bullet(tf,"clean-калибровка > 96 %: мера не «рисует» разрывы сама.",15,color=GREEN_D)
bdata=[("Параметры","random","coherent","learnable","clean"),
       ("1.12B","55.5","49.5","43.8","98.1"),("2.16B","54.5","50.7","43.1","96.4"),
       ("3.08B","56.6","49.3","43.7","99.0")]
bt=s.shapes.add_table(4,5,Inches(5.25),Inches(2.4),Inches(7.45),Inches(2.5)).table
style_table(bt)
for w,c in zip((1.75,1.45,1.5,1.5,1.25),range(5)): bt.columns[c].width=Inches(w)
for j,h in enumerate(bdata[0]): set_cell(bt.cell(0,j),h,14,WHITE,bold=True,fill=GRAPHITE)
tint={1:BAND,2:WHITE,3:REDBG,4:LIGHT}
for i in range(1,4):
    for j,v in enumerate(bdata[i]):
        col=INK;bold=(j==0)
        if j==1: col=GREEN_D;bold=True
        if j==3: col=RED;bold=True
        set_cell(bt.cell(i,j),v,14,col,bold=bold,fill=WHITE if j==0 else tint.get(j,WHITE))
_,tc=box(s,Inches(5.25),Inches(5.2),Inches(7.45),Inches(0.6))
para(tc,"family share, length-matched, null 50.  4.39B не влезает в один H100.",12,MUTED,first=True)
footer(s,10)

# =====================================================================
# 11 — BACK TO THE HOOK
# =====================================================================
s=slide(); header(s,"Ответ на хук","Так почему модель не верит в астрологию?")
rect(s,Inches(0.6),Inches(1.85),Inches(12.1),Inches(1.5),BAND)
_,t=box(s,Inches(0.9),Inches(1.95),Inches(11.5),Inches(1.35),anchor=MSO_ANCHOR.MIDDLE)
para(t,"Задеплоенная модель не верит. Но base-модель после претрейна — почти верит.",19,INK,bold=True,first=True,sa=4)
para(t,"Скептицизм добавляет АЛАЙНМЕНТ (слой 1). Сам объектив претрейна истину не отслеживает "
        "и спокойно моделирует связную псевдонауку (слой 2).",14,GRAPHITE)
_,tf=box(s,Inches(0.6),Inches(3.65),Inches(6.0),Inches(3.2))
para(tf,"Что даёт наша работа",14,GREEN_D,bold=True,first=True,sa=7)
bullet(tf,"Объектив претрейна выбирает выучиваемую, согласованную, сжимаемую "
          "систему — НЕ истинную (контролируемый результат).",14)
bullet(tf,"Поэтому база спокойно моделирует связную псевдонауку → низкий share "
          "на слайде «Слой 2».",14)
bullet(tf,"«Компас истины» у задеплоенных моделей — слой алайнмента, а не "
          "свойство самого объектива.",14)
rect(s,Inches(6.85),Inches(3.65),Inches(5.85),Inches(3.2),AMBBG)
_,tw=box(s,Inches(7.1),Inches(3.8),Inches(5.35),Inches(2.95))
para(tw,"⚠  Обратная сторона — предупреждение",14,AMBER,bold=True,first=True,sa=7)
para(tw,"Где ложь ВЫУЧИВАЕМА и согласованна (наши coherent и learnable) — объектив "
        "предпочитает её, даже правде.",14,INK,sa=7)
para(tw,"Хорошо сконструированная псевдонаука или согласованный бэкдор в данных "
        "переживут фильтрацию по разнообразию.",14,INK)
footer(s,11)

# =====================================================================
# 12 — INSIGHT
# =====================================================================
s=slide(); header(s,"Инсайт","У обучающего объектива нет «компаса истины»")
rect(s,Inches(0.6),Inches(1.85),Inches(12.1),Inches(1.1),BAND)
_,t=box(s,Inches(0.9),Inches(1.9),Inches(11.5),Inches(1.0),anchor=MSO_ANCHOR.MIDDLE)
para(t,"Объектив награждает выучиваемую, согласованную систему — а не истинную.",20,INK,bold=True,first=True)
_,tf=box(s,Inches(0.6),Inches(3.2),Inches(12.1),Inches(3.4))
bullet(tf,"Курация данных: скоординированные/систематические ошибки — именно тот "
          "класс, что это предпочтение усиливает; diversity-фильтры их не ловят.",16,first=True)
bullet(tf,"Отравление данных: согласованный бэкдор выучивается и держит массу лучше "
          "разрозненного шума.",16)
bullet(tf,"Галлюцинации как феномен статистики корпуса (Kalai et al. 2024) — согласуется.",16)
bullet(tf,"Оговорка: все модели ниже продакшн-масштаба; меряем предсказуемость, "
          "«сжатие» оставляем интерпретацией. Хук — мотивация, не доказательство.",14.5,color=MUTED,it=True)
footer(s,12)

# =====================================================================
# 13 — STATUS & NEXT
# =====================================================================
s=slide(); header(s,"Статус и дальше","Где мы и что дальше")
_,tf=box(s,Inches(0.6),Inches(2.0),Inches(6.0),Inches(4.4))
para(tf,"Статус",15,GREEN_D,bold=True,first=True,sa=7)
bullet(tf,"NeurIPS 2026, submission 473 (4/4/3/2); rebuttal + Official Comment поданы.",14.5,mark="•")
bullet(tf,"Мета-ревью требовало одно: показать, что эффект — не mass splitting. "
          "Измерено; часть выжила, часть отозвана.",14.5,mark="•")
bullet(tf,"Ревизия статьи обновлена: learnable-контроль, миллиарды, новые выводы.",14.5,mark="•")
_,tf2=box(s,Inches(6.85),Inches(2.0),Inches(5.85),Inches(4.4))
para(tf2,"Дальше",15,GREEN_D,bold=True,first=True,sa=7)
bullet(tf2,"Низкий конец масштаба (<420M) — контролируемых данных пока нет.",14.5,mark="•")
bullet(tf2,"Единичный realized support при варьируемой предсказуемости.",14.5,mark="•")
bullet(tf2,"Перенос на естественный текст с перечислимым ложным семейством.",14.5,mark="•")
bullet(tf2,"Расширить эмпирический хук: больше моделей, базовые vs aligned.",14.5,mark="•")
bullet(tf2,"Camera-ready: тримминг до лимита страниц, arXiv v2.",14.5,mark="•")
footer(s,13)

out="/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth.pptx"
prs.save(out)
print("saved", out, "| slides:", len(prs.slides._sldIdLst))
