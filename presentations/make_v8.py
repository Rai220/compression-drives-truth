# -*- coding: utf-8 -*-
"""v7 -> v8: предпосылка о противоречивости корпуса (запрос создателя).

  1. Новый слайд «Интернет утверждает всё — и обратное всему» после хука
     (клон карточного макета гипотез): астрология ↔ астрофизика, гомеопатия ↔
     мета-анализы, спорные темы; вывод — «правды данных» не существует,
     модель вынуждена выбирать. Иконки перекрашены в янтарные «↔».
  2. Слайд «Корпус: одна задача — два решения» привязан к вебу: «как в
     реальном вебе, корпус противоречит сам себе».
  3. Номера страниц пересчитаны.

v7 не перезаписывается. Старые make_vN.py не перезапускать.
"""
import copy
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from lxml import etree

SRC = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v7.pptx"
OUT = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v8.pptx"

prs = Presentation(SRC)

def replace_runs(slide, mapping):
    hit = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.text in mapping:
                    hit.add(run.text)
                    run.text = mapping[run.text]
    missed = set(mapping) - hit
    if missed:
        raise SystemExit(f"НЕ НАЙДЕНЫ строки: {missed}")

def find_by_text(text):
    return next(s for s in prs.slides if any(
        sh.has_text_frame and sh.text_frame.text == text for sh in s.shapes))

# 1. Клон слайда гипотез -> слайд противоречивости корпуса
donor = find_by_text("Что решает: частота, длина или сжимаемость?")
clone = prs.slides.add_slide(donor.slide_layout)
for ph in list(clone.shapes):
    ph._element.getparent().remove(ph._element)
for sh in donor.shapes:
    clone.shapes._spTree.append(copy.deepcopy(sh._element))
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
clone._element.find(f"{P_NS}cSld").insert(
    0, copy.deepcopy(donor._element.find(f".//{P_NS}bg")))

replace_runs(clone, {
    "ГИПОТЕЗЫ": "ИСХОДНЫЙ ФАКТ",
    "Что решает: частота, длина или сжимаемость?":
        "Интернет утверждает всё — и обратное всему",
    # карточка 1
    "Данных больше": "Астрология",
    "Правильного в вебе просто больше — модель верит большинству.":
        "«Положение звёзд определяет судьбу» — гороскопы, тысячи сайтов.",
    "Нет: астрологии в вебе море, а модели её отвергают. В нашем "
    "эксперименте частота зафиксирована 50/50.":
        "И рядом — астрофизика и научпоп: «никакого влияния нет». "
        "Оба пласта огромны, оба в претрейне.",
    # карточка 2
    "Правда короче": "Гомеопатия",
    "Верные тексты компактнее, а вероятность любит короткое.":
        "«Вылечила всю семью!» — форумы, отзывы, магазины разведений.",
    "Нет — но это главный конфаунд: вероятность любит короткое, "
    "поэтому длина везде жёстко уравнена.":
        "И рядом — мета-анализы: эффект неотличим от плацебо. "
        "Модель читает и то и другое.",
    # карточка 3
    "Правда логична": "Спорные темы",
    "У правды внутренняя логика: вся система описывается кратко.":
        "Диеты, политика, крипта, ИИ: на каждый текст найдётся анти-текст.",
    "Да: модель — почти идеальный архиватор. Её цель — не истина, "
    "а лучшее объяснение наблюдаемого мира.":
        "У корпуса нет единой позиции: тексты с несовместимыми взглядами "
        "лежат в одном претрейне.",
    # иконки и нижняя строка
    "✕": "↔",
    "✓": "↔",
    "Но сначала выясним, чей голос мы вообще слышим в ответах — самой "
    "модели или её воспитания.":
        "Значит, «правды данных» не существует — корпус спорит сам с собой, "
        "и модель вынуждена выбирать. А модели почему-то единодушны. Разбираемся.",
})

# иконки: все кружки и глифы «↔» — в янтарный
AMBER = "E8B23A"
for sh in clone.shapes:
    x = etree.tostring(sh._element).decode()
    is_icon = "ellipse" in x or (sh.has_text_frame and sh.text_frame.text.strip() == "↔")
    if is_icon:
        for clr in sh._element.iter(qn("a:srgbClr")):
            if clr.get("val") in ("FF6B5A", "4FD787"):
                clr.set("val", AMBER)

KICK_TOP, KICK_LEFT = 502920, 566928
width = Emu(402338 + 75895 * len("ИСХОДНЫЙ ФАКТ"))
for sh in clone.shapes:
    if sh.top == KICK_TOP and sh.left == KICK_LEFT:
        sh.width = width

# позиция: сразу после хука «6 моделей»
hook = find_by_text("6 моделей от разных команд — почти одна картина мира")
hook_pos = list(prs.slides).index(hook)
lst = prs.slides._sldIdLst
clone_el = lst[-1]
lst.remove(clone_el)
lst.insert(hook_pos + 1, clone_el)

# 2. Наш корпус = интернет в миниатюре
corpus = find_by_text("Корпус: одна задача — два решения")
replace_runs(corpus, {
    "Каждая задача встречается и с верным, и с неверным решением 50/50 — "
    "частота, формат и домен зафиксированы.":
        "Как в реальном вебе, корпус противоречит сам себе: каждая задача "
        "встречается и с верным, и с неверным решением 50/50 — частота, "
        "формат и домен зафиксированы.",
})

# 3. Номера страниц по позиции
for pos, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t.isdigit() and sh.left and sh.left > Emu(10000000):
                sh.text_frame.paragraphs[0].runs[0].text = str(pos)

prs.save(OUT)
print(f"saved {OUT} | slides: {len(prs.slides._sldIdLst)}")
