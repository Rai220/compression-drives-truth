# -*- coding: utf-8 -*-
"""v5 -> v6: развёрнута история проверки базы (запрос создателя: «какая была
идея, про алайнмент, про гипотезу с претрейном и как я её проверял»).

Идея (сл. 2-3), алайнмент (сл. 7-8) и гипотеза (баннер сл. 7) уже в колоде;
не хватало методики зонда — она была сжата в один абзац на «Слое 2».

  1. Новый слайд «Слой 2 — как спросить» между «Слоем 1» и «Слоем 2»
     (клон двухпанельного «Холодного душа»): панель «Зонд» — как подслушиваем
     вероятности у base-модели; панель «Ловушки» — почему нельзя просто
     спросить (база не чат; наивный зонд меряет длину, частица «не»
     переворачивает знак).
  2. «Слой 2» стал чистым слайдом результата: вводный абзац методики заменён
     ссылкой на зонд, нижняя сноска упрощена.
  3. Номера страниц пересчитаны.

v5 не перезаписывается. Старые make_vN.py не перезапускать.
"""
import copy
from pptx import Presentation
from pptx.util import Emu

SRC = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v5.pptx"
OUT = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v6.pptx"

prs = Presentation(SRC)
slides = list(prs.slides)

def replace_runs(slide, mapping):
    hit = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.text in mapping:
                    hit.add(run.text)
                    run.text = mapping[run.text]
    missed = set(mapping) - hit
    if missed:
        raise SystemExit(f"НЕ НАЙДЕНЫ строки: {missed}")

def find_by_title(text):
    return next(s for s in prs.slides if any(
        sh.has_text_frame and sh.text_frame.text == text for sh in s.shapes))

# ---------------------------------------------------------------
# 1. Клон «Холодного душа» -> методика зонда
# ---------------------------------------------------------------
donor = find_by_title("Рецензент поставил «2» — и оказался прав")
clone = prs.slides.add_slide(donor.slide_layout)
for ph in list(clone.shapes):
    ph._element.getparent().remove(ph._element)
for sh in donor.shapes:
    clone.shapes._spTree.append(copy.deepcopy(sh._element))
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
clone._element.find(f"{P_NS}cSld").insert(
    0, copy.deepcopy(donor._element.find(f".//{P_NS}bg")))

replace_runs(clone, {
    "ХОЛОДНЫЙ ДУШ": "СЛОЙ 2 — КАК СПРОСИТЬ",
    "Рецензент поставил «2» — и оказался прав":
        "База на вопросы не отвечает. Зато вероятности не врут",
    # левая панель (✓): зонд
    "ВЫЖИЛО": "ЗОНД",
    "86M: верное держит 59.0 % ± 0.8 массы против random (≈ +9 пп).":
        "Даём утверждение и две концовки равной длины: «Это утверждение "
        "верно / ложно», «Ответ: да / нет».",
    "Coherent НЕ фильтруется ни на каком размере — робастная половина.":
        "share = P(научный вердикт) / (P(научный) + P(псевдо)). "
        "50 — безразличие, выше — наклон к науке.",
    "Реплика на независимом корпусе: 59.4 %.":
        "Qwen2.5 base 0.5B / 1.5B / 3B — претрейн без SFT и RLHF.",
    # правая панель (✕): ловушки
    "ОТОЗВАНО": "ЛОВУШКИ",
    "Multi-rule: pair acc 47→79, а масса 47.9→47.8. Это mass splitting.":
        "База — не чат: инструкции игнорирует, вопрос просто продолжает "
        "как текст. «Спросить» нельзя.",
    "Description-length ladder: масса плоская (55.6 / 51.9 / 52.6 / 51.3).":
        "Наивный зонд меряет длину, а не мнение: одна частица «не» "
        "переворачивает знак.",
    "«Coherent ниже случайности» — это длина завершения.":
        "Поэтому концовки строго равной длины — и контрольные темы, "
        "где ответ заранее известен.",
    "Возражение (mass splitting) я померил честно: половина выводов умерла. "
    "Выжившая половина оказалась интереснее убитой.":
        "Калибровка: на темах консенсуса (эволюция) зонд даёт 77–84 — "
        "наклон он ловит. Теперь — псевдонаука.",
})

KICK_TOP, KICK_LEFT = 502920, 566928
width = Emu(402338 + 75895 * len("СЛОЙ 2 — КАК СПРОСИТЬ"))
for sh in clone.shapes:
    if sh.top == KICK_TOP and sh.left == KICK_LEFT:
        sh.width = width

layer2 = find_by_title("У base-модели без RLHF нет компаса истины")
layer2_pos = list(prs.slides).index(layer2)
lst = prs.slides._sldIdLst
clone_el = lst[-1]
lst.remove(clone_el)
lst.insert(layer2_pos, clone_el)

# ---------------------------------------------------------------
# 2. «Слой 2» -> чистый результат
# ---------------------------------------------------------------
replace_runs(layer2, {
    "Base-модель не следует инструкциям — «спросить» нельзя. Меряем "
    "правдоподобием: P(научный вердикт) против P(псевдонаучного), длина "
    "контролируется.":
        "Зонд с прошлого слайда, семь тем: пять псевдонаучных и две "
        "контрольные, где консенсус «за».",
    "Base Qwen2.5, длина ответов уравнена. В наивной пробе знак share "
    "двигает даже частица «не» — тот же length-confound, что вскрывает "
    "статья.":
        "Base Qwen2.5, два шаблона на тему, длина концовок уравнена.",
})

# ---------------------------------------------------------------
# 3. Номера страниц по позиции
# ---------------------------------------------------------------
for pos, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t.isdigit() and sh.left and sh.left > Emu(10000000):
                sh.text_frame.paragraphs[0].runs[0].text = str(pos)

prs.save(OUT)
print(f"saved {OUT} | slides: {len(prs.slides._sldIdLst)}")
