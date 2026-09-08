# -*- coding: utf-8 -*-
"""v3 -> v4: поверх ручных правок создателя (титул «Сжатие приводит к истине?»,
слайд Кеплера, AI RnD Day).

  1. Новый слайд «Как вообще учат модель» (клон «Механизма»): претрейн → SFT →
     RL-шейпинг; вопрос «а есть ли предпочтения у голой базы?» — мостик к
     слоям 1/2. Ставится между «Гипотезами» и «Слоем 1».
  2. Футер обновлён под новый титул и площадку: «Сжатие приводит к истине? · AI RnD Day».
  3. Номера страниц пересчитаны по фактической позиции (после вставки Кеплера
     они съехали на единицу).

v3 не перезаписывается. НЕ пересобирать из make_v3.py — там нет правок создателя.
"""
import copy
from pptx import Presentation
from pptx.util import Emu

SRC = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v3.pptx"
OUT = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v4.pptx"

prs = Presentation(SRC)
slides = list(prs.slides)

def replace_runs(slide, mapping):
    hit = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.text in mapping:
                    hit.add(run.text)
                    run.text = mapping[run.text]
    missed = set(mapping) - hit
    if missed:
        raise SystemExit(f"НЕ НАЙДЕНЫ строки на слайде: {missed}")

# ---------------------------------------------------------------
# 1. Клон «Механизма» -> «Как вообще учат модель»
# ---------------------------------------------------------------
mech = next(s for s in slides if any(
    sh.has_text_frame and sh.text_frame.text.startswith("Почему знак переворачивается")
    for sh in s.shapes))
clone = prs.slides.add_slide(mech.slide_layout)
for ph in list(clone.shapes):
    ph._element.getparent().remove(ph._element)
for sh in mech.shapes:
    clone.shapes._spTree.append(copy.deepcopy(sh._element))
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
src_bg = mech._element.find(f".//{P_NS}bg")
clone._element.find(f"{P_NS}cSld").insert(0, copy.deepcopy(src_bg))

replace_runs(clone, {
    "МЕХАНИЗМ": "КАК ВООБЩЕ УЧАТ МОДЕЛЬ",
    "Почему знак переворачивается — и почему это не артефакт":
        "Претрейн строит картину мира, алайнмент её причёсывает",
    "Выученный селектор": "Претрейн",
    "концентрирует ложную массу на одной предсказуемой строке":
        "Триллионы токенов веба, одна задача: предскажи следующий "
        "токен. Здесь и возникает картина мира.",
    "Модель «коммитится»": "SFT",
    "realized share ложной ветви растёт до 66 → 89 → 92":
        "Примеры «как надо отвечать»: форма диалога, роль ассистента, "
        "стиль — какими их выбрали создатели.",
    "Знак переворачивается": "RL и другой шейпинг",
    "family share падает ниже null: 43 → 49 → 49":
        "Награда за ответы, которые нравятся людям и проходят политику. "
        "Дружный скепсис к псевдонауке — отсюда.",
    "Почему below-null — это предпочтение, а не некомпетентность?":
        "Задеплоенная модель отвечает так, как её воспитали. А сама база?",
    "random и learnable обучены на ОДНИХ верных выводах — правильный ответ "
    "считают одинаково. random выше нуля, learnable ниже, при равной "
    "компетентности. Причина — именно селектор. ":
        "Таблица «Нет / Спорно» из начала — ответы ПОСЛЕ алайнмента: это "
        "голос воспитания. Собственные предпочтения модели надо искать до "
        "него — у голой базы, сразу после претрейна. ",
    "Это та causal identification, которую раньше «не знали, как построить».":
        "Есть ли они у неё вообще? Снимем слои по одному и посмотрим.",
    "learnable_mod детерминирован, но НЕ выучивается → читается как random. "
    "Важна реальная выучиваемость, а не детерминизм.":
        "Сначала — рука алайнмента (слой 1), затем — что осталось под ним "
        "(слой 2).",
})

# пилюля кикера: ширина линейна от длины текста (замерено по v2)
KICK_TOP, KICK_LEFT = 502920, 566928
width = Emu(402338 + 75895 * len("КАК ВООБЩЕ УЧАТ МОДЕЛЬ"))
for sh in clone.shapes:
    if sh.top == KICK_TOP and sh.left == KICK_LEFT:
        sh.width = width

# позиция: сразу после «Гипотез», перед «Слой 1 — алайнмент»
layer1_pos = next(i for i, s in enumerate(prs.slides) if any(
    sh.has_text_frame and "СЛОЙ 1" in sh.text_frame.text for sh in s.shapes))
lst = prs.slides._sldIdLst
clone_el = lst[-1]
lst.remove(clone_el)
lst.insert(layer1_pos, clone_el)

# ---------------------------------------------------------------
# 2. Футер под новый титул; 3. номера страниц по позиции
# ---------------------------------------------------------------
OLD_FOOT = "Сжатие предпочитает согласованность, а не истину  ·  GigaRnD"
NEW_FOOT = "Сжатие приводит к истине?  ·  AI RnD Day"
for pos, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == OLD_FOOT:
            sh.text_frame.paragraphs[0].runs[0].text = NEW_FOOT
        elif t.isdigit() and sh.left and sh.left > Emu(10000000):
            sh.text_frame.paragraphs[0].runs[0].text = str(pos)

prs.save(OUT)
print(f"saved {OUT} | slides: {len(prs.slides._sldIdLst)}")
