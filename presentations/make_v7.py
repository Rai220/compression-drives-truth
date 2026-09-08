# -*- coding: utf-8 -*-
"""v6 -> v7: из истории поиска — в итог (запрос создателя: убрать рецензента и
ошибочные гипотезы, рассказывать «что мы сейчас знаем»).

  1. Удалён слайд «Рецензент поставил «2» — и оказался прав» (выжило/отозвано).
  2. Титул отвечает на собственный вопрос: «Короткий ответ: нет...».
  3. Слайд идеи: без «сел доказывать» и «название не пережило».
  4. Гипотезы: вердикты «нет / нет / да» вместо детектива; третьей карточке —
     зелёная галка (была янтарная ≈).
  5. Метрика: «Pair accuracy обманывает — меряем массу» вместо «как я обманул
     сам себя»; learnable — «ГЛАВНЫЙ РЕЗУЛЬТАТ»; масштаб и механизм без
     «раньше было/не знали»; статус без ребаттла.

v6 не перезаписывается. Старые make_vN.py не перезапускать.
"""
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from lxml import etree

SRC = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v6.pptx"
OUT = "/Users/krestnikov/giga/compression-drives-truth/presentations/gigarnd_compression_truth_v7.pptx"

prs = Presentation(SRC)

def replace_runs(slide, mapping):
    hit = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.text in mapping:
                    hit.add(run.text)
                    run.text = mapping[run.text]
    missed = set(mapping) - hit
    if missed:
        raise SystemExit(f"НЕ НАЙДЕНЫ строки: {missed}")

def find_by_text(text):
    return next(s for s in prs.slides if any(
        sh.has_text_frame and sh.text_frame.text == text for sh in s.shapes))

KICK_TOP, KICK_LEFT = 502920, 566928
def resize_kicker(slide, text):
    width = Emu(402338 + 75895 * len(text))
    for sh in slide.shapes:
        if sh.top == KICK_TOP and sh.left == KICK_LEFT:
            sh.width = width

# 1. Удалить слайд про рецензента
victim = find_by_text("Рецензент поставил «2» — и оказался прав")
lst = prs.slides._sldIdLst
for sldId in list(lst):
    rId = sldId.get(qn("r:id"))
    if prs.part.rels[rId].target_part is victim.part:
        prs.part.drop_rel(rId)
        lst.remove(sldId)
        break

# 2. Титул: вопрос заголовка получает ответ
replace_runs(prs.slides[0], {
    "Я хотел доказать, что предсказание следующего токена обречено "
    "сходиться к истине. Доказал обратное. Рассказываю, как так вышло.":
        "Короткий ответ: нет. Модель предпочитает ту картину мира, которая "
        "лучше всего описывает данные, — даже если она ложна.",
})

# 3. Слайд идеи: без биографии
idea = find_by_text("Красивая гипотеза: сжатие должно любить истину")
replace_runs(idea, {
    "Красиво? Мне тоже понравилось — настолько, что я назвал репозиторий "
    "compression-drives-truth и сел доказывать неизбежность истины. ":
        "Красивая гипотеза — и проверяемая: дальше устроим ей очную ставку "
        "с данными. ",
    "Спойлер: название эксперименты не пережило.":
        "Спойлер: всё окажется тоньше — и интереснее.",
    "Март 2026, план: «Compression Truth Bias — индуктивное смещение LLM "
    "к истине». Но сначала — во что вообще верят готовые модели?":
        "Но сначала — во что вообще верят готовые модели?",
})

# 4. Гипотезы: вердикты-итоги, третьей — зелёную галку
hyp = find_by_text("Три подозреваемых: частота, длина, сжимаемость")
replace_runs(hyp, {
    "Три подозреваемых: частота, длина, сжимаемость":
        "Что решает: частота, длина или сжимаемость?",
    "Парадокс выше это уже убил: астрологии море, а модели её отвергают. "
    "В эксперименте частота будет ровно 50/50.":
        "Нет: астрологии в вебе море, а модели её отвергают. В нашем "
        "эксперименте частота зафиксирована 50/50.",
    "Коварная гипотеза: длину придётся жёстко уравнивать. Дважды чуть "
    "не обманула меня самого — увидите.":
        "Нет — но это главный конфаунд: вероятность любит короткое, "
        "поэтому длина везде жёстко уравнена.",
    "Главный подозреваемый: модель — почти идеальный архиватор. "
    "Тогда её цель — не истина, а лучшее объяснение наблюдаемого "
    "мира.":
        "Да: модель — почти идеальный архиватор. Её цель — не истина, "
        "а лучшее объяснение наблюдаемого мира.",
    "≈": "✓",
})
GREEN = "4FD787"
for sh in hyp.shapes:
    x = etree.tostring(sh._element).decode()
    if ("E8B23A" in x) and (sh.has_text_frame is False or "ellipse" in x
                            or sh.text_frame.text.strip() == "✓"):
        for clr in sh._element.iter(qn("a:srgbClr")):
            if clr.get("val") == "E8B23A":
                clr.set("val", GREEN)

# 5. Метрика, learnable, масштаб, механизм, статус
metric = find_by_text("Первые цифры были слишком красивыми")
replace_runs(metric, {
    "КАК Я ОБМАНУЛ САМ СЕБЯ": "МЕТРИКА",
    "Первые цифры были слишком красивыми":
        "Pair accuracy обманывает — меряем массу",
})
resize_kicker(metric, "МЕТРИКА")

learn = find_by_text("Сделай ложь выучиваемой — и модель предпочтёт её истине")
replace_runs(learn, {"НОВЫЙ РЕЗУЛЬТАТ · ГЛАВНОЕ": "ГЛАВНЫЙ РЕЗУЛЬТАТ"})
resize_kicker(learn, "ГЛАВНЫЙ РЕЗУЛЬТАТ")

scale = find_by_text("Контраст держится до настоящих 3.08B")
replace_runs(scale, {
    "Раньше потолок был 490M. Теперь — Qwen3 на 1.12B / 2.16B / 3.08B.":
        "Qwen3-архитектура: 1.12B / 2.16B / 3.08B, обучение с нуля.",
})

mech = find_by_text("Почему знак переворачивается — и почему это не артефакт")
replace_runs(mech, {
    "Это та causal identification, которую раньше «не знали, как построить».":
        "Чистая causal identification: между условиями меняется только селектор.",
})

status = find_by_text("Где мы и что дальше")
replace_runs(status, {
    "Ключевое возражение — mass splitting — измерено напрямую. Часть "
    "результатов выжила, часть отозвана честно.":
        "Препринт: arXiv 2603.11749; статья на рецензии NeurIPS 2026.",
})

# 6. Номера страниц по позиции
for pos, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t.isdigit() and sh.left and sh.left > Emu(10000000):
                sh.text_frame.paragraphs[0].runs[0].text = str(pos)

prs.save(OUT)
print(f"saved {OUT} | slides: {len(prs.slides._sldIdLst)}")
